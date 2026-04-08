def 将病史txt转换成doc(x):
    import re
    from docx import Document
    from docx.shared import Pt
    from docx.oxml.ns import qn
    
    with open(x,'r',encoding='utf-8') as f:
        病史集=f.read().strip()[2:-2]      #把json文件开头的["和末尾的"]删去
        病史集=re.sub(r'(\\n|\\t)+','\n',病史集).strip()      #把中间的换行符和制表符全部替换成换行；不过json文件中的\n和\t被解析为字符串后会被表达为\\n和\\t；但是我们希望用真正的换行去替掉，所以替换要用\n
        病史集=re.split(r'","',病史集)          #字符串中的","原本在json文件中就是两份病史的分隔处
    #读取txt文件

    doc=Document()
    style=doc.styles['Normal']        #设置字体时，这行是基础
    style.font.name='Times New Roman'    #英文和数字用这个
    style._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')   #中文要单独更改
    style.font.size=Pt(12)   #统一字号

    for i in 病史集:
        doc.add_paragraph(i)
        doc.add_paragraph('\n')
        doc.add_paragraph('\n')
        doc.add_paragraph('\n')
        doc.add_paragraph('\n')
        doc.add_paragraph('\n')
        doc.add_paragraph('\n')
        
    doc.save(x[:-4]+'docx')
    print('病史文件已生成')
