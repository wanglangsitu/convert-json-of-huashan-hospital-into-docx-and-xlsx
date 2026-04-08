from tkinter import filedialog
path=filedialog.askopenfilename()
with open(path,'r',encoding='utf-8') as f:
    examines=eval(f.read().strip())
#读取txt文件

from openpyxl import Workbook
import re
document_new=Workbook()
last_row={}
for i in range(len(examines)):
    examine=examines[i]
    sheet_name=examine['key'].strip()
    sheet_name=re.sub(r'[\\/*?:\[\]]','',sheet_name)   # \/?:*[]这七个字符不能出现在excel的sheetname中
    if examine['value'][0]!='检查结果未出':
        if sheet_name not in last_row:        
            sheet_new=document_new.create_sheet(title=sheet_name)        
            for j in range(len(examine['value'])):
                for k in range(len(examine['value'][j])):
                    sheet_new.cell(row=j+1,column=k+1,value=examine['value'][j][k])
            last_row.update({sheet_name:len(examine['value'])})
        else:
            sheet_new=document_new[sheet_name]
            first_row=last_row[sheet_name]+3
            for j in range(len(examine['value'])):
                for k in range(len(examine['value'][j])):
                    sheet_new.cell(row=j+first_row,column=k+1,value=examine['value'][j][k])
            last_row[sheet_name]=first_row+len(examine['value'])-1
#将同种检查合并到一张sheet内，每次检查之间空两行
document_new.save(path[:-5]+'.xlsx')
#生成检验excel文档

print('文件已生成')


    
