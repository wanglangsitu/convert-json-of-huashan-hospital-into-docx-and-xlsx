def 将医嘱txt转换成excel(x):
    from openpyxl import Workbook
    
    with open(x,'r',encoding='utf-8') as f:
        医嘱集=eval(f.read().strip())
    #读取txt文件

    document_new=Workbook()
    sheet_new=document_new.create_sheet(title='药物医嘱')
    sheet_new.cell(row=1,column=1,value='药物名')
    sheet_new.cell(row=1,column=2,value='医嘱时限')
    sheet_new.cell(row=1,column=3,value='用法')
    sheet_new.cell(row=1,column=4,value='剂量')
    sheet_new.cell(row=1,column=5,value='单位')
    sheet_new.cell(row=1,column=6,value='频次')
    sheet_new.cell(row=1,column=7,value='开始日期')
    sheet_new.cell(row=1,column=8,value='结束日期')

    for i in range(len(医嘱集)):
        医嘱=医嘱集[i]
        
        try:
            医嘱时限=医嘱[0]
            开始日期=医嘱[2]
            结束日期=医嘱[3]
            其余信息=医嘱[1].split(' ')
            药物名=其余信息[0]
            用法=其余信息[1][3:]
            剂量=其余信息[2][3:]
            单位=其余信息[3][3:]
            频次=其余信息[5][3:]
            
            sheet_new.cell(row=i+2,column=2,value=医嘱时限)
            sheet_new.cell(row=i+2,column=7,value=开始日期)
            sheet_new.cell(row=i+2,column=8,value=结束日期)
            sheet_new.cell(row=i+2,column=1,value=药物名)
            sheet_new.cell(row=i+2,column=3,value=用法)
            sheet_new.cell(row=i+2,column=4,value=剂量)
            sheet_new.cell(row=i+2,column=5,value=单位)        
            sheet_new.cell(row=i+2,column=6,value=频次)
        except:
            pass

    document_new.save(x[:-4]+'xlsx')
    print('医嘱文件已生成')